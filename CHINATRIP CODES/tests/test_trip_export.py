"""行程导出测试（2026-10-23，gap #2「行程的出口」）。

两层：
- 渲染器单测：手工构造 ItineraryPlan 直接喂三种渲染器（不经过校验器，
  断言与铁律规则解耦）；
- API 测试：app 组装照抄 tests/test_itineraries_api.py（fresh store +
  dependency_overrides），覆盖三种格式 / 非法格式 / 跨 owner / 缺设备头。
"""

import io
from datetime import time as dtime

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.itineraries import (
    get_entitlement_service,
    get_store,
)
from app.api.itineraries import (
    router as itineraries_router,
)
from app.services.entitlement import EntitlementService, InMemoryQuotaStore
from app.services.itinerary.models import (
    DayPlan,
    GeoPoint,
    ItineraryPlan,
    TransferItem,
    VisitItem,
)
from app.services.itinerary_store import InMemoryItineraryStore
from app.services.trip_export import (
    render_docx,
    render_html,
    render_xlsx,
    safe_filename,
)

DEV_A = {"X-Device-ID": "device-a"}
DEV_B = {"X-Device-ID": "device-b"}


def make_plan() -> ItineraryPlan:
    visit = VisitItem(
        item_id="v1",
        name="故宫",
        city="北京",
        poi_id="poi_1",
        location=GeoPoint(lon=116.397, lat=39.918),
        address="北京市东城区景山前街4号",
        opening_hours="08:30-17:00",
        day=1,
        start_time=dtime(9, 0),
        duration_minutes=120,
        distance_to_next_km=1.2,
        next_walk_minutes=15,
        leg_source="amap",
        pinned=True,
    )
    transfer = TransferItem(
        item_id="t1",
        day=2,
        seq=1,
        from_city="北京",
        to_city="西安",
        mode="train",
        start_time=dtime(8, 0),
        duration_minutes=270,
        train_no="G87",
    )
    weird = VisitItem(
        item_id="v2",
        name="奇<em>怪</em>名&B",
        city="西安",
        day=2,
        start_time=dtime(13, 30),
        duration_minutes=90,
    )
    return ItineraryPlan(
        itinerary_id="it_test1",
        cities=["北京", "西安"],
        days=[
            DayPlan(day=1, base_city="北京", items=[visit]),
            DayPlan(day=2, base_city="西安", items=[transfer, weird]),
        ],
        assumptions=["假设1：未计入排队时间"],
    )


# --- 渲染器单测 -------------------------------------------------------------


def test_render_html_contains_days_and_escapes():
    text = render_html(make_plan())
    assert "第 1 天 · 北京" in text
    assert "第 2 天 · 西安" in text
    assert "故宫" in text and "北京市东城区景山前街4号" in text
    assert "08:30-17:00" in text
    assert "已钉住" in text and "步行约 15 分钟" in text and "高德实测" in text
    assert "火车/高铁 北京 → 西安" in text and "G87" in text
    assert "假设1：未计入排队时间" in text
    # XSS 防护（TREK escHtml 口径）：文本一律转义
    assert "奇&lt;em&gt;怪&lt;/em&gt;名&amp;B" in text
    assert "<em>" not in text and "<script" not in text


def test_render_docx_contains_days_and_details():
    doc = Document(io.BytesIO(render_docx(make_plan())))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "行程单：北京 → 西安" in text
    assert "第 1 天 · 北京" in text
    assert "故宫" in text and "地址：北京市东城区景山前街4号" in text
    assert "开放时间：08:30-17:00" in text and "已钉住" in text
    assert "火车/高铁 北京 → 西安（G87）" in text
    assert "假设1：未计入排队时间" in text


def test_render_xlsx_contains_days_and_notes():
    ws = load_workbook(io.BytesIO(render_xlsx(make_plan()))).active
    assert "行程单：北京 → 西安" in str(ws["A1"].value)
    assert ws["A2"].value == "时间"
    flat = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
    assert any("第 1 天 · 北京" in v for v in flat)
    assert "故宫" in flat
    assert any("已钉住" in v and "距下一站 1.2 km" in v for v in flat)
    assert any("火车/高铁 北京 → 西安" in v for v in flat)


def test_safe_filename_strips_illegal_chars():
    plan = make_plan()
    plan.itinerary_id = "it/x?1"
    name = safe_filename(plan)
    assert "/" not in name and "?" not in name
    assert "北京-西安-行程" in name


# --- API --------------------------------------------------------------------


def make_app() -> TestClient:
    app = FastAPI()
    store = InMemoryItineraryStore()  # 先建实例再绑定，保证同 app 内共享
    app.dependency_overrides[get_store] = lambda: store
    app.dependency_overrides[get_entitlement_service] = lambda: EntitlementService(
        InMemoryQuotaStore(), rules=[]
    )
    app.include_router(itineraries_router)
    return TestClient(app)


def create_itinerary(client: TestClient, headers=DEV_A) -> str:
    resp = client.post(
        "/api/v1/itineraries", json={"cities": ["北京"], "days": 2}, headers=headers
    )
    assert resp.status_code == 201
    return resp.json()["itinerary_id"]


def test_export_three_formats_on_empty_plan():
    client = make_app()
    iid = create_itinerary(client)
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=pdf", headers=DEV_A)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/html")
    assert "行程单：北京" in resp.text
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=docx", headers=DEV_A)
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"  # docx 是 zip 容器
    assert "wordprocessingml" in resp.headers["content-type"]
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=xlsx", headers=DEV_A)
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "spreadsheetml" in resp.headers["content-type"]


def test_export_content_disposition_utf8_filename():
    client = make_app()
    iid = create_itinerary(client)
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=xlsx", headers=DEV_A)
    cd = resp.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd and ".xlsx" in cd


def test_export_bad_format_answers_422():
    client = make_app()
    iid = create_itinerary(client)
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=csv", headers=DEV_A)
    assert resp.status_code == 422


def test_export_foreign_owner_answers_404():
    client = make_app()
    iid = create_itinerary(client)
    resp = client.get(f"/api/v1/itineraries/{iid}/export?format=pdf", headers=DEV_B)
    assert resp.status_code == 404


def test_export_missing_device_answers_400():
    client = make_app()
    iid = create_itinerary(client, headers=DEV_A)
    assert client.get(f"/api/v1/itineraries/{iid}/export?format=pdf").status_code == 400
