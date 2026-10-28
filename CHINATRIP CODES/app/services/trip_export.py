"""行程导出渲染器（2026-10-23，gap #2「行程的出口」）。

三种格式、两种生成路径：
- pdf：打印优化的 HTML 字符串，前端放进沙箱 iframe 调浏览器打印
  （「另存为 PDF」）——liketrek/TREK 同款零依赖方案，中文矢量文字零风险；
- docx / xlsx：服务端生成二进制（python-docx / openpyxl）。

纯函数、无 DB 无网络；输入一律是 ItineraryPlan，输出 str | bytes。
"""

import html
import io
import re
from datetime import time

from app.services.itinerary.models import ItineraryPlan, TransferItem
from app.services.lang_ctx import L, is_en


def _mode_label(mode: str) -> str:
    return {
        "train": L("火车/高铁", "Train/High-speed rail"),
        "flight": L("飞机", "Flight"),
        "car": L("驾车", "Drive"),
        "bus": L("大巴", "Bus"),
    }.get(mode, mode)


def _pin_label() -> str:
    return L("📌 已钉住", "📌 Pinned")


def _paren(inner: str) -> str:
    """成对括号：zh 全角，EN 半角。"""
    return f"（{inner}）" if not is_en() else f"({inner})"


def plan_title(plan: ItineraryPlan) -> str:
    return L("行程单", "Itinerary") + L("：", ": ") + " → ".join(plan.cities)


def safe_filename(plan: ItineraryPlan) -> str:
    """下载文件名：城市路由 + 行程编号，去掉 Windows 非法字符。"""
    raw = "-".join(plan.cities) + L("-行程-", "-itinerary-") + plan.itinerary_id
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "", raw)
    return cleaned[:80] or L("行程", "Itinerary")


def _fmt_time(t: time | None) -> str:
    return t.strftime("%H:%M") if t else ""


def _visit_note(
    *,
    pinned: bool,
    distance_km: float | None,
    walk_min: int | None,
    leg_source: str | None,
) -> str:
    notes = []
    if pinned:
        notes.append(_pin_label())
    if isinstance(distance_km, (int, float)):
        seg = L(
            f"距下一站 {distance_km:.1f} km",
            f"{distance_km:.1f} km to next stop",
        )
        if walk_min is not None:
            seg += L(f" · 步行约 {walk_min} 分钟", f" · about {walk_min} min walk")
        seg += (
            L("（地图实测）", " (Map measured)")
            if leg_source in ("amap", "baidu")
            else L("（直线估算）", " (straight-line estimate)")
        )
        notes.append(seg)
    return " · ".join(notes)


def display_rows(plan: ItineraryPlan) -> list[dict]:
    """按天拍平成展示行，三种渲染器共用同一份口径（与前端 planCard 一致）。"""
    days = []
    for d in plan.days:
        rows = []
        for it in d.items:
            if isinstance(it, TransferItem):
                mode = _mode_label(it.mode)
                rows.append(
                    {
                        "kind": "transfer",
                        "time": _fmt_time(it.start_time),
                        "title": f"{mode} {it.from_city} → {it.to_city}",
                        "city": "",
                        "address": it.train_no or "",
                        "hours": "",
                        "minutes": it.duration_minutes,
                        "note": "",
                    }
                )
            else:
                rows.append(
                    {
                        "kind": "visit",
                        "time": _fmt_time(it.start_time),
                        "title": it.name,
                        "city": it.city,
                        "address": it.address or "",
                        "hours": it.opening_hours or "",
                        "minutes": it.duration_minutes,
                        "note": _visit_note(
                            pinned=it.pinned,
                            distance_km=it.distance_to_next_km,
                            walk_min=it.next_walk_minutes,
                            leg_source=it.leg_source,
                        ),
                    }
                )
        days.append({"day": d.day, "base_city": d.base_city, "rows": rows})
    return days


def visit_count(plan: ItineraryPlan) -> int:
    return sum(
        1 for d in plan.days for it in d.items if not isinstance(it, TransferItem)
    )


def _meta_line(plan: ItineraryPlan) -> str:
    return (
        L(f"共 {len(plan.days)} 天", f"{len(plan.days)} days")
        + " · "
        + L(f"{visit_count(plan)} 个站点", f"{visit_count(plan)} stops")
    )


# --- PDF：打印 HTML -------------------------------------------------------------


def render_html(plan: ItineraryPlan) -> str:
    esc = html.escape
    days_html = []
    for d in display_rows(plan):
        body = []
        for r in d["rows"]:
            if r["kind"] == "transfer":
                detail = esc(r["address"]) if r["address"] else ""
                cells = (
                    f'<td class="t">{esc(r["time"]) or "—"}</td>'
                    f"<td>{esc(r['title'])}</td>"
                    f'<td colspan="2" class="sub">{detail}</td>'
                )
                body.append(f'<tr class="transfer">{cells}</tr>')
                continue
            sub = []
            if r["minutes"]:
                sub.append(L(f"约 {r['minutes']} 分钟", f"About {r['minutes']} min"))
            if r["note"]:
                sub.append(esc(r["note"]))
            sub_html = f'<div class="sub">{" · ".join(sub)}</div>' if sub else ""
            badge = ""
            if _pin_label() in r["note"]:
                badge = f' <span class="badge">{L("已钉住", "Pinned")}</span>'
            body.append(
                f'<tr><td class="t">{esc(r["time"]) or "—"}</td>'
                f"<td>{esc(r['title'])}{badge}{_paren(esc(r['city']))}{sub_html}</td>"
                f"<td>{esc(r['address']) or '—'}</td>"
                f"<td>{esc(r['hours']) or '—'}</td></tr>"
            )
        empty_row = (
            '<tr><td colspan="4" class="sub">'
            f"{L('（本日暂无安排）', '(No plans for this day)')}</td></tr>"
        )
        body_html = "".join(body) or empty_row
        day_head = (
            L(f"第 {d['day']} 天", f"Day {d['day']}") + f" · {esc(d['base_city'])}"
        )
        days_html.append(
            f'<div class="day"><h2>{day_head}</h2>'
            "<table><thead><tr>"
            f"<th>{L('时间', 'Time')}</th><th>{L('安排', 'Itinerary')}</th>"
            f"<th>{L('地址', 'Address')}</th><th>{L('开放时间', 'Opening hours')}</th>"
            "</tr></thead>"
            f"<tbody>{body_html}</tbody></table></div>"
        )
    assumptions = ""
    if plan.assumptions:
        items = "".join(f"<li>{esc(a)}</li>" for a in plan.assumptions)
        assumptions = (
            f'<h2>{L("假设标注", "Assumptions")}</h2><ul class="assume">{items}</ul>'
        )
    page_lang = "en" if is_en() else "zh-CN"
    return (
        f'<!doctype html><html lang="{page_lang}"><head><meta charset="utf-8">'
        f"<title>{esc(plan_title(plan))}</title><style>"
        'body{font-family:"Microsoft YaHei","PingFang SC",system-ui,sans-serif;'
        "margin:0;padding:28px;color:#1a1a1a;font-size:13px}"
        "h1{font-size:20px;margin:0 0 6px}"
        ".meta{color:#555;margin-bottom:10px}"
        ".day{margin-top:22px;padding-top:12px;"
        "border-top:2px solid #333;page-break-after:always}"
        ".day:last-of-type{page-break-after:auto}"
        "h2{font-size:15px;margin:0 0 8px}"
        "table{width:100%;border-collapse:collapse}"
        "th,td{border:1px solid #c9c9c9;"
        "padding:5px 8px;text-align:left;vertical-align:top}"
        "th{background:#f2f2f2}"
        ".t{white-space:nowrap;width:56px;color:#555}"
        "tr.transfer td{background:#fafafa;color:#444}"
        ".badge{font-size:10px;border:1px solid #b8860b;"
        "color:#8a6d00;border-radius:3px;padding:0 3px}"
        ".sub{color:#666;font-size:11px;margin-top:2px}"
        "ul.assume{margin:4px 0;padding-left:20px;color:#444}"
        "footer{margin-top:26px;color:#999;font-size:10px;text-align:center}"
        "@media print{.no-print{display:none}body{padding:0}}"
        "@media screen{body{max-width:760px;margin:0 auto}"
        ".no-print{display:block;background:#fffbe6;"
        "border:1px solid #e6d87a;border-radius:6px;padding:8px 12px;"
        "margin-bottom:14px;font-size:12px;color:#6b5900}}"
        "</style></head><body>"
        '<div class="no-print">'
        + L(
            "💡 打印对话框中请选择「另存为 PDF」（目标打印机），即可得到 PDF 文件。",
            '💡 In the print dialog, choose "Save as PDF" as the '
            "destination to get a PDF file.",
        )
        + "</div>"
        f"<h1>{esc(plan_title(plan))}</h1>"
        f'<div class="meta">{L("行程编号", "Itinerary ID")} '
        f"<code>{esc(plan.itinerary_id)}</code> · {_meta_line(plan)}</div>"
        f"{''.join(days_html)}{assumptions}"
        f"<footer>{L('由 ChinaTrip 生成', 'Generated by ChinaTrip')}"
        "</footer></body></html>"
    )


# --- Word ----------------------------------------------------------------------


def render_docx(plan: ItineraryPlan) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def eastasia(style, name: str = "微软雅黑") -> None:
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), name)

    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1"):
        try:
            eastasia(doc.styles[style_name])
        except KeyError:
            pass
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    run = title.add_run(plan_title(plan))
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"{L('行程编号', 'Itinerary ID')} {plan.itinerary_id} · {_meta_line(plan)}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for d in display_rows(plan):
        doc.add_heading(
            L(f"第 {d['day']} 天", f"Day {d['day']}") + f" · {d['base_city']}", level=1
        )
        for r in d["rows"]:
            if r["kind"] == "transfer":
                extra = _paren(r["address"]) if r["address"] else ""
                doc.add_paragraph(f"{r['time'] or '—'}  {r['title']}{extra}")
                continue
            head = doc.add_paragraph()
            bits = [r["time"] or "—", r["title"], _paren(r["city"])]
            run = head.add_run("  ".join(bits))
            run.bold = True
            if _pin_label() in r["note"]:
                head.add_run(f"  {_pin_label()}")
            details = []
            if r["address"]:
                details.append(L("地址：", "Address: ") + r["address"])
            if r["hours"]:
                details.append(L("开放时间：", "Opening hours: ") + r["hours"])
            if r["minutes"]:
                details.append(
                    L(f"约 {r['minutes']} 分钟", f"About {r['minutes']} min")
                )
            rest_note = " · ".join(
                n for n in r["note"].split(" · ") if _pin_label() not in n
            )
            if rest_note:
                details.append(rest_note)
            if details:
                sub = doc.add_paragraph("\n".join(details))
                sub.paragraph_format.left_indent = Pt(18)
                sub.runs[0].font.size = Pt(9)
                sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if plan.assumptions:
        doc.add_heading(L("假设标注", "Assumptions"), level=1)
        for a in plan.assumptions:
            doc.add_paragraph(a, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Excel ---------------------------------------------------------------------


def render_xlsx(plan: ItineraryPlan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = L("行程", "Itinerary")

    header_fill = PatternFill("solid", fgColor="333333")
    day_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    title = f"{plan_title(plan)}{_paren(plan.itinerary_id)}"
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)

    headers = [
        L("时间", "Time"),
        L("安排", "Itinerary"),
        L("城市", "City"),
        L("地址", "Address"),
        L("开放时间", "Opening hours"),
        L("时长(分钟)", "Duration (min)"),
        L("备注", "Notes"),
    ]
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill

    row = 3
    for d in display_rows(plan):
        label = L(f"第 {d['day']} 天", f"Day {d['day']}") + f" · {d['base_city']}"
        ws.cell(row=row, column=1, value=label).font = bold
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = day_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1
        for r in d["rows"]:
            ws.cell(row=row, column=1, value=r["time"])
            ws.cell(row=row, column=2, value=r["title"]).alignment = wrap
            ws.cell(row=row, column=3, value=r["city"])
            ws.cell(row=row, column=4, value=r["address"]).alignment = wrap
            ws.cell(row=row, column=5, value=r["hours"]).alignment = wrap
            ws.cell(row=row, column=6, value=r["minutes"])
            ws.cell(row=row, column=7, value=r["note"]).alignment = wrap
            row += 1

    widths = (8, 36, 10, 34, 18, 10, 24)
    for col, width in zip(range(1, 8), widths, strict=True):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
